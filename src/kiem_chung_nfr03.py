"""Kiểm chứng NFR-03 – precision/recall của phát hiện cặp nghi trùng (FR-04, src/etl/match.py). TD-32.

Chạy từ thư mục gốc repo (đã kích hoạt .venv) – KHÔNG cần PostgreSQL:

    # Dữ liệu mẫu trong repo: nhãn đúng lấy từ script sinh mẫu (script biết chắc hồ sơ nào là cùng một khách)
    python -m src.kiem_chung_nfr03 --tep data/mau/customers_mau.csv --nhan-dung data/mau/nhan_dung_mau.csv

    # Dữ liệu thật (LMS, không commit): (1) tạo phiếu 300 cặp → (2) em gán nhãn tay cột "nhan" → (3) tính
    python -m src.kiem_chung_nfr03 --tep data/raw/customers_raw.csv --tao-phieu data/nhan/nfr03_phieu_gan_nhan.xlsx
    python -m src.kiem_chung_nfr03 --tep data/raw/customers_raw.csv --nhan data/nhan/nfr03_phieu_gan_nhan.xlsx

Khung cặp dựng lại từ tệp nguồn bằng đúng các hàm thuần của luồng chính (extract → transform → match), nên phép
kiểm chứng không phụ thuộc trạng thái CSDL và không ghi gì vào CSDL; cùng tệp + seed 42 → cùng 300 cặp, cùng
kết quả (NFR-05). Tệp gọi pha E và pha T, không gọi pha L (tests/test_ranh_gioi_lop.py).
Mã thoát: 0 đã tính xong (in ĐẠT / CHƯA ĐẠT) · 1 lỗi kỹ thuật · 2 tệp nguồn hoặc phiếu nhãn bị từ chối.
"""
import argparse
import pathlib
import sys
from datetime import datetime

import pandas as pd

from src import config
from src.etl import danh_gia_match as dg
from src.etl import extract, match, transform

MA_THANH_CONG, MA_LOI_KY_THUAT, MA_TU_CHOI = 0, 1, 2
COT_HO_SO = {'ho_ten': 'full_name_raw', 'so_dien_thoai': 'phone_raw', 'email': 'email_raw',
             'dia_chi': 'address_raw', 'ngay_tao': 'created_raw'}


class TuChoi(Exception):
    """Thao tác bị từ chối (mã thoát 2)."""


def so(n) -> str:
    return f'{int(round(n)):,}'.replace(',', '.')


def pt(x) -> str:
    return '–' if x is None else f'{x * 100:.2f}%'.replace('.', ',')


# ------------------------------------------------------------------ dựng khung từ tệp nguồn (pha E + pha T)
def doc_ho_so(tep) -> pd.DataFrame:
    """Tệp nguồn → bảng hồ sơ [ma (số dòng), source_record_id, *_raw, phone_norm, name_key].
    Giống src/pipeline.py: name_key tính từ name_norm; hồ sơ có mã lỗi không tham gia so khớp (QL7-04) nên để
    trống phone_norm của chúng trong khung."""
    bang = transform.chuan_hoa(extract.sang_bang_tam(extract.doc_tep(tep)))
    bang['name_key'] = bang['name_norm'].map(match.tao_name_key)
    bang['phone_norm'] = bang['phone_norm'].where(bang['error_code'].isna(), None)
    return bang.rename(columns={'row_no': 'ma'})


def lap_mau(ho_so: pd.DataFrame, nguong: float, seed: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    khung = dg.khung_cap(ho_so[['ma', 'phone_norm', 'name_key']], nguong)
    return khung, dg.chon_mau(khung, dg.CO_MAU, seed)


# ------------------------------------------------------------------ phiếu gán nhãn
def bang_phieu(mau: pd.DataFrame, ho_so: pd.DataFrame) -> pd.DataFrame:
    """Phiếu MÙ: chỉ có giá trị gốc của hai hồ sơ, không có tầng, điểm hay quyết định của hệ thống – để nhãn
    là phán đoán độc lập của người, không bị điểm số dẫn dắt."""
    hs = ho_so.set_index('ma')
    phieu = pd.DataFrame({'ma_cap': mau['ma_cap']})
    for ben in ('a', 'b'):
        ma = mau[f'ma_{ben}']
        phieu[f'dong_{ben}'] = ma.to_numpy()
        for ten, cot in COT_HO_SO.items():
            phieu[f'{ten}_{ben}'] = hs.loc[ma, cot].to_numpy()
    phieu['nhan'] = ''
    phieu['ghi_chu'] = ''
    return phieu


HUONG_DAN = [
    'PHIẾU GÁN NHÃN NFR-03 – mỗi dòng là một cặp hồ sơ khách hàng (hồ sơ A ở cột *_a, hồ sơ B ở cột *_b).',
    'Cột nhan: 1 = hai hồ sơ là CÙNG MỘT khách hàng · 0 = hai khách hàng KHÁC nhau. Không để trống.',
    'Xem đủ các cột: họ tên, số điện thoại, email, địa chỉ, ngày tạo – không đoán theo thứ tự dòng.',
    'Cặp khó quyết định: vẫn chọn 1 hoặc 0, ghi lý do ngắn vào cột ghi_chu.',
    'Không sửa các cột khác; không thêm, xóa hay sắp xếp lại dòng.',
    'Lưu lại đúng tên tệp (.xlsx), rồi chạy: python -m src.kiem_chung_nfr03 --tep <tệp nguồn> --nhan <tệp này>',
]


def ghi_phieu(phieu: pd.DataFrame, duong_dan: pathlib.Path) -> None:
    """Ghi .xlsx (mọi ô dạng văn bản để Excel/Numbers không làm mất số 0 đầu) hoặc .csv UTF-8 có BOM."""
    if duong_dan.exists():
        raise TuChoi(f'{duong_dan} đã có – không ghi đè vì có thể đã chứa nhãn em gán; đổi tên tệp cũ nếu muốn tạo lại')
    duong_dan.parent.mkdir(parents=True, exist_ok=True)
    if duong_dan.suffix.lower() == '.csv':
        phieu.to_csv(duong_dan, index=False, encoding='utf-8-sig')
        return
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    with pd.ExcelWriter(duong_dan, engine='openpyxl') as w:
        phieu.astype(str).to_excel(w, sheet_name='phieu', index=False)
        pd.DataFrame({'huong_dan': HUONG_DAN}).to_excel(w, sheet_name='huong_dan', index=False)
        ws = w.sheets['phieu']
        ws.freeze_panes = 'B2'
        for o in ws[1]:
            o.font = Font(bold=True)
        cot_nhan = phieu.columns.get_loc('nhan') + 1
        for dong in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for o in dong:
                o.number_format = '@'
            dong[cot_nhan - 1].fill = PatternFill('solid', fgColor='FFF2CC')
        dv = DataValidation(type='list', formula1='"1,0"', allow_blank=True,
                            error='Chỉ nhập 1 (cùng một khách) hoặc 0 (khác khách)')
        ws.add_data_validation(dv)
        chu = ws.cell(row=1, column=cot_nhan).column_letter
        dv.add(f'{chu}2:{chu}{ws.max_row}')
        for c in ws.columns:
            ws.column_dimensions[c[0].column_letter].width = min(34, max(8, *(len(str(o.value or '')) + 2
                                                                           for o in c[:60])))
        w.sheets['huong_dan'].column_dimensions['A'].width = 120


def doc_phieu(duong_dan: pathlib.Path) -> pd.DataFrame:
    if not duong_dan.is_file():
        raise TuChoi(f'Không tìm thấy phiếu nhãn {duong_dan}')
    if duong_dan.suffix.lower() == '.xlsx':
        return pd.read_excel(duong_dan, sheet_name=0, dtype=str, keep_default_na=False)
    try:
        return pd.read_csv(duong_dan, dtype=str, keep_default_na=False, encoding='utf-8-sig')
    except UnicodeDecodeError:
        raise TuChoi(f'Phiếu "{duong_dan.name}" không phải UTF-8 – lưu lại dạng CSV UTF-8 hoặc dùng .xlsx') from None


def nhan_tu_tep_dung(mau: pd.DataFrame, ho_so: pd.DataFrame, tep_dung: pathlib.Path) -> pd.DataFrame:
    """Dữ liệu mẫu: nhãn = hai hồ sơ có cùng ma_khach trong data/mau/nhan_dung_mau.csv (do script sinh mẫu ghi)."""
    dung = pd.read_csv(tep_dung, dtype=str, keep_default_na=False)
    khach = ho_so.set_index('ma')['source_record_id'].map(dict(zip(dung['record_id'], dung['ma_khach'])))
    if khach.isna().any():
        raise TuChoi(f'{tep_dung.name} không có ma_khach cho {int(khach.isna().sum())} hồ sơ – không phải tệp nhãn '
                     'đúng của tệp nguồn này')
    return pd.DataFrame({'ma_cap': mau['ma_cap'], 'ma_a': mau['ma_a'], 'ma_b': mau['ma_b'],
                         'nhan': (khach.loc[mau['ma_a']].to_numpy() == khach.loc[mau['ma_b']].to_numpy()).astype(int)})


# ------------------------------------------------------------------ in kết quả
def in_tang(khung, mau, in_ra) -> None:
    in_ra(f'      {"Tầng":<5} {"Trong khung":>11} {"Lấy mẫu":>8}   Mô tả')
    for d in dg.dem_theo_tang(khung, mau).itertuples():
        in_ra(f'      {d.tang:<5} {so(d.trong_khung):>11} {d.lay_mau:>8}   {d.mo_ta}')


def in_ket_qua(bang: pd.DataFrame, nguong: float, in_ra) -> dict:
    kq = dg.tinh_chi_so(bang, nguong)
    dau = {True: '✔ đạt', False: '✘ chưa đạt'}
    in_ra(f'      Ma trận nhầm lẫn          nhãn: cùng một khách    nhãn: khác khách')
    in_ra(f'        hệ thống GHÉP               TP = {kq["tp"]:<8}           FP = {kq["fp"]}')
    in_ra(f'        hệ thống KHÔNG ghép         FN = {kq["fn"]:<8}           TN = {kq["tn"]}')
    in_ra(f'      Precision = TP/(TP+FP) = {kq["tp"]}/{kq["tp"] + kq["fp"]} = {pt(kq["precision"])}'
          f'  (yêu cầu ≥ {dg.YEU_CAU_PRECISION}%)  {dau[kq["dat_precision"]]}')
    in_ra(f'      Recall    = TP/(TP+FN) = {kq["tp"]}/{kq["tp"] + kq["fn"]} = {pt(kq["recall"])}'
          f'  (yêu cầu ≥ {dg.YEU_CAU_RECALL}%)  {dau[kq["dat_recall"]]}')
    in_ra(f'      → NFR-03 ở ngưỡng {nguong:g}: {"ĐẠT" if kq["dat"] else "CHƯA ĐẠT"}')
    return kq


def ghi_bao_cao(nguon: str, tep, nguong, khung, mau, bang, kq, quet, uoc, thu_muc=None) -> pathlib.Path:
    thu_muc = thu_muc or config.duong_dan(config.doc_cau_hinh(), 'thu_muc_bao_cao')
    d = [f'# NFR-03 – precision/recall phát hiện trùng · {datetime.now():%d/%m/%Y %H:%M}', '',
         f'- Tệp nguồn: `{tep}` · ngưỡng {nguong:g} · nguồn nhãn: {nguon}',
         f'- Khung {so(len(khung))} cặp · tập nhãn {len(bang)} cặp '
         f'({int(bang["nhan"].sum())} cùng một khách, {int((bang["nhan"] == 0).sum())} khác khách)', '',
         '| Tầng | Trong khung | Lấy mẫu | Mô tả |', '|---|---:|---:|---|']
    d += [f'| {r.tang} | {so(r.trong_khung)} | {r.lay_mau} | {r.mo_ta} |' for r in dg.dem_theo_tang(khung, mau).itertuples()]
    d += ['', f'TP {kq["tp"]} · FP {kq["fp"]} · FN {kq["fn"]} · TN {kq["tn"]} → precision {pt(kq["precision"])}, '
          f'recall {pt(kq["recall"])} → **{"ĐẠT" if kq["dat"] else "CHƯA ĐẠT"}**', '',
          '| Ngưỡng | TP | FP | FN | Precision | Recall | NFR-03 |', '|---:|---:|---:|---:|---:|---:|---|']
    d += [f'| {r.nguong:g} | {r.tp} | {r.fp} | {r.fn} | {pt(r.precision)} | {pt(r.recall)} | '
          f'{"đạt" if r.dat else "chưa đạt"} |' for r in quet.itertuples()]
    d += ['', f'Ước lượng cho toàn khung (trọng số theo tầng): precision {pt(uoc["precision"])}, '
          f'recall {pt(uoc["recall"])}', '']
    p = thu_muc / f'nfr03-{pathlib.Path(tep).stem}-{datetime.now():%Y%m%d_%H%M%S}.md'
    p.write_text('\n'.join(d) + '\n', encoding='utf-8')
    return p


def chay(tep, nguong: float, seed: int, tao_phieu=None, nhan=None, nhan_dung=None, in_ra=print,
         thu_muc_bao_cao=None) -> dict:
    in_ra(f'NFR-03 – kiểm chứng phát hiện cặp nghi trùng · tệp {tep} · ngưỡng {nguong:g} · seed {seed}')
    ho_so = doc_ho_so(tep)
    khung, mau = lap_mau(ho_so, nguong, seed)
    in_ra(f'[1/3] khung cặp  {so(len(ho_so))} hồ sơ → {so(len(khung))} cặp trong khung → '
          f'{len(mau)} cặp để gán nhãn' + (' (khung ≤ 300 cặp: lấy hết)' if len(mau) == len(khung) else ''))
    in_tang(khung, mau, in_ra)

    if tao_phieu is not None:
        p = pathlib.Path(tao_phieu)
        ghi_phieu(bang_phieu(mau, ho_so), p)
        in_ra(f'[2/3] phiếu gán nhãn (không có điểm, tầng hay quyết định của hệ thống): {p}')
        in_ra('      Điền cột nhan: 1 = cùng một khách · 0 = hai khách khác nhau (hướng dẫn ở trang "huong_dan").')
        in_ra(f'[3/3] gán xong thì chạy: python -m src.kiem_chung_nfr03 --tep {tep} --nhan {p}')
        return {'khung': khung, 'mau': mau}

    if nhan_dung is not None:
        nguon = f'nhãn đúng của script sinh mẫu ({pathlib.Path(nhan_dung).name})'
        bang_nhan = nhan_tu_tep_dung(mau, ho_so, pathlib.Path(nhan_dung))
    else:
        nguon = f'nhãn trên phiếu ({pathlib.Path(nhan).name})'
        bang_nhan = dg.kiem_phieu_nhan(doc_phieu(pathlib.Path(nhan)))
    bang = dg.gan_nhan_vao_khung(khung, bang_nhan, mau)
    in_ra(f'[2/3] nhãn       {len(bang)} cặp – {nguon}: {int(bang["nhan"].sum())} cùng một khách · '
          f'{int((bang["nhan"] == 0).sum())} khác khách')
    in_ra(f'[3/3] kết quả trên {len(bang)} cặp có nhãn, ngưỡng {nguong:g} (QL7-02):')
    kq = in_ket_qua(bang, nguong, in_ra)
    quet = dg.quet_nguong(bang)
    in_ra('      Quét ngưỡng (chỉnh ngưỡng nếu chưa đạt):')
    in_ra(f'        {"Ngưỡng":>6} {"TP":>5} {"FP":>5} {"FN":>5} {"Precision":>10} {"Recall":>9}  NFR-03')
    for r in quet.itertuples():
        in_ra(f'        {r.nguong:>6g} {r.tp:>5} {r.fp:>5} {r.fn:>5} {pt(r.precision):>10} {pt(r.recall):>9}  '
              f'{"đạt" if r.dat else "chưa đạt"}')
    uoc = dg.tinh_chi_so(bang, nguong, co_trong_so=True)
    if len(mau) < len(khung):
        in_ra(f'      Ước lượng cho toàn khung {so(len(khung))} cặp (trọng số theo tầng): precision '
              f'{pt(uoc["precision"])} · recall {pt(uoc["recall"])}')
    p = ghi_bao_cao(nguon, tep, nguong, khung, mau, bang, kq, quet, uoc, thu_muc_bao_cao)
    in_ra(f'Báo cáo: {p.relative_to(config.GOC) if p.is_relative_to(config.GOC) else p}')
    return {'khung': khung, 'mau': mau, 'bang': bang, 'ket_qua': kq, 'quet': quet, 'uoc_luong': uoc}


def main(argv=None) -> int:
    cau_hinh = config.doc_cau_hinh()
    ap = argparse.ArgumentParser(description='NFR-03: precision/recall của phát hiện cặp nghi trùng (không cần CSDL)')
    ap.add_argument('--tep', required=True, help='Tệp khách hàng nguồn (.csv/.xlsx) – cùng tệp đã nạp ở pipeline')
    viec = ap.add_mutually_exclusive_group(required=True)
    viec.add_argument('--tao-phieu', help='Ghi phiếu 300 cặp để gán nhãn (.xlsx hoặc .csv)')
    viec.add_argument('--nhan', help='Phiếu đã gán nhãn (cột nhan) hoặc tệp nhãn .csv → tính precision/recall')
    viec.add_argument('--nhan-dung', help='Tệp nhãn đúng record_id → ma_khach (dữ liệu mẫu)')
    ap.add_argument('--nguong', type=float, default=float(cau_hinh['nguong_fuzzy']), help='Ngưỡng FUZZY (QL7-02)')
    a = ap.parse_args(argv)
    try:
        chay(a.tep, a.nguong, int(cau_hinh['seed']), a.tao_phieu, a.nhan, a.nhan_dung)
        return MA_THANH_CONG
    except (TuChoi, extract.TepBiTuChoi, dg.NhanKhongHopLe) as e:
        print(f'TỪ CHỐI: {e}', file=sys.stderr)
        return MA_TU_CHOI
    except Exception as e:                           # noqa: BLE001 – báo lỗi gọn cho người dùng
        print(f'LỖI: {type(e).__name__}: {e}', file=sys.stderr)
        return MA_LOI_KY_THUAT


if __name__ == '__main__':
    sys.exit(main())
